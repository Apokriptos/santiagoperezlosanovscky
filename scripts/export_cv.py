from pathlib import Path
from xml.sax.saxutils import escape
import re
from docx import Document
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from pypdf import PdfReader

repo=Path(__file__).resolve().parents[1]
source=repo/'documentation'/'CV — Technical Game Designer — Santiago Perez Losanovscky.docx'
target=repo/'documentation'/'cv-technical-game-designer.pdf'
doc=Document(source)
pdfmetrics.registerFont(TTFont('Arial',r'C:\Windows\Fonts\arial.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Bold',r'C:\Windows\Fonts\arialbd.ttf'))
pdfmetrics.registerFont(TTFont('Arial-Italic',r'C:\Windows\Fonts\ariali.ttf'))
pdfmetrics.registerFont(TTFont('Arial-BoldItalic',r'C:\Windows\Fonts\arialbi.ttf'))
pdfmetrics.registerFontFamily('Arial',normal='Arial',bold='Arial-Bold',italic='Arial-Italic',boldItalic='Arial-BoldItalic')
section=doc.sections[0]
default_size=doc.styles['Normal'].font.size.pt if doc.styles['Normal'].font.size else 9
story=[]
for index,p in enumerate(doc.paragraphs):
    if not p.text.strip():
        # Preserve explicit source page breaks without adding empty paragraphs.
        if p._p.xpath('.//w:br[@w:type="page"]'): story.append(PageBreak())
        continue
    fmt=p.paragraph_format
    size=next((r.font.size.pt for r in p.runs if r.font.size),default_size)
    is_heading=p.text.isupper() and len(p.text)<60
    if p.text=='PROFESSIONAL EXPERIENCE (Continued)': story.append(PageBreak())
    style=ParagraphStyle('source',fontName='Arial',fontSize=size,leading=size*1.13,
        spaceBefore=fmt.space_before.pt if fmt.space_before else 0,
        spaceAfter=fmt.space_after.pt if fmt.space_after else 1.2,
        keepWithNext=bool(is_heading or (p.runs and p.runs[0].bold and size>=9.5)),
        splitLongWords=False)
    text=''
    for run in p.runs:
        value=escape(run.text).replace('\n','<br/>')
        if run.bold: value='<b>'+value+'</b>'
        if run.italic: value='<i>'+value+'</i>'
        text+=value
    story.append(Paragraph(text,style))
SimpleDocTemplate(str(target),pagesize=(section.page_width/12700,section.page_height/12700),
    leftMargin=section.left_margin/12700,rightMargin=section.right_margin/12700,
    topMargin=section.top_margin/12700,bottomMargin=section.bottom_margin/12700,
    title='Santiago Pérez Losanovscky — Technical Game Designer CV',author='Santiago Pérez Losanovscky').build(story)
reader=PdfReader(target)
normal=lambda s: re.sub(r'\s+','',s)
pdf_text=normal(''.join(page.extract_text() for page in reader.pages))
for p in doc.paragraphs:
    if p.text.strip(): assert normal(p.text) in pdf_text, 'Missing source text: '+p.text[:50]
print('Exported',target.name,'pages:',len(reader.pages),'— every source paragraph verified in PDF text')

# The HTML companion is a real searchable document, not a raster preview.
# It works in embedded browsers that do not expose a native PDF plugin.
html_parts=[]
for index,p in enumerate(doc.paragraphs):
    if not p.text.strip(): continue
    tag='h1' if index==0 else 'h2' if (p.text.isupper() and len(p.text)<60) or p.text=='PROFESSIONAL EXPERIENCE (Continued)' else 'p'
    text=''
    for run in p.runs:
        value=escape(run.text).replace('\n','<br>')
        if run.bold:value='<strong>'+value+'</strong>'
        if run.italic:value='<em>'+value+'</em>'
        text+=value
    html_parts.append(f'<{tag}>{text}</{tag}>')
html='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Santiago Pérez Losanovscky CV</title><style>
*{box-sizing:border-box}body{margin:0;padding:28px;background:#e5e1d9;color:#171717;font:13px/1.6 Arial,sans-serif;overflow-wrap:anywhere}
main{max-width:850px;margin:auto;padding:32px;background:#fff;box-shadow:0 2px 10px #0001}h1{font-size:23px;line-height:1.2;margin:0 0 8px}h2{font-size:15px;line-height:1.3;margin:25px 0 10px}p{margin:0 0 10px}
@media(max-width:450px){body{padding:12px}main{padding:20px}h1{font-size:21px}}@media print{body,main{padding:0;background:#fff;box-shadow:none}h2{break-after:avoid}p{orphans:3;widows:3}}
</style><main>'''+''.join(html_parts)+'</main></html>'
preview=target.with_suffix('.html');preview.write_text(html,encoding='utf-8')
print('Exported searchable HTML companion:',preview.name)

